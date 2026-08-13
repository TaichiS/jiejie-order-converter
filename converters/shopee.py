"""
converters/shopee.py
蝦皮訂單轉換器。
輸入：Order.toship.YYYYMMDD_YYYYMMDD.xlsx（工作表 orders）
輸出：{MMDD}蝦_new.xlsx
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from openpyxl import Workbook

from converters.base import BaseConverter, ConversionResult, RowError, open_xlsx, SHOPEE_XLSX_PASSWORD
from app.utils import output_path


# 欄位索引常數
COL_ORDER_ID    = 0
COL_PRODUCT     = 24   # 完整商品名稱
COL_PROD_ID     = 25   # 商品ID（品號）
COL_OPTION      = 26   # 商品選項名稱（短品名）
COL_ORIG_PRICE  = 29   # 商品原價
COL_ACT_PRICE   = 30   # 商品活動價格
COL_QTY         = 33   # 數量
COL_SHIPPING_FEE= 7    # 買家支付運費
COL_COIN        = 12   # 蝦幣折抵
COL_COUPON      = 17   # 優惠券
COL_SHIP_DATE   = 49   # 最晚出貨日期


class ShopeeConverter(BaseConverter):

    @property
    def source_type(self) -> str:
        return "shopee"

    def _validate(self, input_files: list[Path]) -> list[RowError]:
        errors = []
        xlsx = [f for f in input_files if f.suffix.lower() == ".xlsx"]
        if not xlsx:
            errors.append(RowError(0, "file", "", "找不到 .xlsx 訂單檔"))
            return errors

        try:
            wb = open_xlsx(xlsx[0], password=SHOPEE_XLSX_PASSWORD)
            if "orders" not in wb.sheetnames:
                errors.append(RowError(0, "sheet", xlsx[0].name, "工作表 'orders' 不存在",
                                       source_file=xlsx[0].name))
        except Exception as e:
            errors.append(RowError(0, "file", xlsx[0].name, f"無法開啟：{e}",
                                   source_file=xlsx[0].name))
        return errors

    def _process(self, input_files: list[Path]) -> ConversionResult:
        order_file = next(f for f in input_files if f.suffix.lower() == ".xlsx")

        wb_in  = open_xlsx(order_file, password=SHOPEE_XLSX_PASSWORD)
        ws_in  = wb_in["orders"]
        rows   = list(ws_in.iter_rows(values_only=True))
        header = list(rows[0])

        # 從第一列資料取出出貨日期 YYYYMMDD
        order_date = ""
        if len(rows) > 1 and rows[1][COL_SHIP_DATE]:
            raw_date = str(rows[1][COL_SHIP_DATE])[:10]  # "2026-03-30"
            order_date = raw_date.replace("-", "")        # "20260330"

        # MMDD 供輸出檔名
        mmdd = order_date[4:8] if len(order_date) == 8 else "0000"

        # 建立輸出活頁簿
        wb_out = Workbook()
        ws_out = wb_out.active
        ws_out.title = "orders"
        ws_out.append(header)

        success_count = 0
        fail_count    = 0
        manual_count  = 0
        errors: list[RowError] = []
        seen_orders: set[str] = set()

        for row_idx, raw_row in enumerate(rows[1:], start=2):
            row = list(raw_row)
            # 跳過全空列（read_only 模式可能多讀 Excel dimension 尾端空列）
            if row[COL_ORDER_ID] is None and row[COL_PRODUCT] is None:
                continue
            order_id     = str(row[COL_ORDER_ID]) if row[COL_ORDER_ID] else ""
            full_name    = str(row[COL_PRODUCT])  if row[COL_PRODUCT]  else ""
            option_name  = str(row[COL_OPTION])   if row[COL_OPTION]   else ""
            qty_order    = row[COL_QTY] or 0

            # 1. 蝦皮同一商品頁可能有多個選項；選項名稱才是實際下單品項。
            #    先查選項，無法命中時才回退到商品名稱。
            act_price = float(row[COL_ACT_PRICE]) if row[COL_ACT_PRICE] else 0
            lookup_names = _build_lookup_names(option_name, full_name)
            product = None
            candidates = []
            for short_name in lookup_names:
                product, candidates = self._lookup_by_name(short_name, amount=act_price)
                if product:
                    break
            if not product:
                searched = "、".join(f"「{name}」" for name in lookup_names)
                errors.append(RowError(
                    row_idx, "商品名稱／選項名稱", full_name,
                    f"找不到 {searched}",
                    candidates=candidates,
                    source_file=order_file.name,
                ))
                fail_count += 1
                continue

            # 3. 轉換數值
            pack_count       = int(product.get("包數") or 1)
            checkout_price   = float(product.get("商品結帳價") or 0)
            new_qty          = int(qty_order) * pack_count

            row[COL_PROD_ID]   = product["品號"]
            row[COL_OPTION]    = product["品名"]
            row[COL_ORIG_PRICE]= act_price           # 份數定價（原始訂單值）
            row[COL_ACT_PRICE] = checkout_price      # 每包單價
            row[COL_QTY]       = new_qty

            # 4. 優惠券欄位全數清零（折扣只顯示 M 欄蝦幣折抵）
            row[COL_COUPON] = 0

            # 5. 同訂單第二列起費用清零
            if order_id in seen_orders:
                row[COL_COIN]        = 0
                row[COL_SHIPPING_FEE]= 0
            else:
                seen_orders.add(order_id)

            ws_out.append(row)
            success_count += 1

        # 儲存輸出
        out_file = output_path(
            self.output_dir.parent,
            self.source_type,
            _parse_date(order_date),
            f"{mmdd}蝦_new.xlsx",
        )
        wb_out.save(out_file)

        status = "completed" if fail_count == 0 else ("partial" if success_count > 0 else "failed")
        return ConversionResult(
            source_type   = self.source_type,
            success_count = success_count,
            fail_count    = fail_count,
            manual_count  = manual_count,
            order_date    = order_date,
            output_files  = [out_file],
            errors        = errors,
            status        = status,
        )


def _extract_short_name(full_name: str) -> str:
    """
    從完整商品名稱中提取短品名。
    例：「【捷捷寶寶粥】2-1 醬燒什錦豬肉 | 冷凍副食品 ...」→ 「2-1醬燒什錦豬肉」
    """
    # 移除品牌前綴 【xxx】
    name = re.sub(r"^【[^】]*】\s*", "", full_name)
    # 取 | 之前的部分
    name = name.split("|")[0]
    # 移除所有空白
    name = re.sub(r"\s+", "", name)
    return name.strip()


def _build_lookup_names(option_name: str, full_name: str) -> list[str]:
    """建立蝦皮品項查詢順序：實際選項優先，商品主標題作為備援。"""
    names = []
    for raw_name in (option_name, full_name):
        name = _extract_short_name(raw_name)
        if name and name not in {"-", "無規格"} and name not in names:
            names.append(name)
    return names


def _parse_date(order_date: str):
    from datetime import datetime
    try:
        return datetime.strptime(order_date, "%Y%m%d")
    except (ValueError, TypeError):
        return datetime.today()

"""
converters/xuantu.py
炫兔團購訂單轉換器。
輸入：炫兔後台匯出 xlsx（預設工作表，含「訂單號碼」「商品貨號」欄）
輸出：MMDD{配送方式}.xlsx（16 欄標準格式）

轉換邏輯：
- 跳過商品貨號為空的贈品列
- 數量 × 包數、價格 ÷ 包數
- 移除商品名稱前綴「團購限定B|」
- 加購品類型 → 加購折扣；全家服務編號 / 7-11 店號 → 送貨編號
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook

from converters.base import BaseConverter, ConversionResult, RowError
from app.utils import output_path

OUTPUT_HEADER = [
    "訂單號碼", "收件人", "完整地址", "收件人電話號碼", "發票號碼",
    "商品貨號", "商品名稱", "數量", "商品結帳價", "商品折扣優惠",
    "商品折扣金額", "加購折扣", "點數折現分攤", "出貨備註", "送貨編號", "付款方式",
]

_DATE_SLASH = re.compile(r'(\d{4})/(\d{2})/(\d{2})')
_CODE_NORM  = re.compile(r'^(\d+[A-Z]?)-([A-Z]?)0*(\d+)$')


def _normalize_code(code: str) -> str:
    """去掉短代碼中的前置零：'2-01' → '2-1'，'0-17' → '0-17'（無前置零不變）"""
    m = _CODE_NORM.match(code.strip())
    if m:
        return m.group(1) + '-' + m.group(2) + str(int(m.group(3)))
    return code.strip()


def _delivery_type(filename: str) -> str:
    """從檔名提取配送方式（全家/黑貓），找不到回傳空字串"""
    if '全家' in filename:
        return '全家'
    if '黑貓' in filename:
        return '黑貓'
    return ''


class XuantuConverter(BaseConverter):

    @property
    def source_type(self) -> str:
        return "xuantu"

    def _validate(self, input_files: list[Path]) -> list[RowError]:
        xlsx = [f for f in input_files if f.suffix.lower() == ".xlsx"]
        if not xlsx:
            return [RowError(0, "file", "", "找不到 .xlsx 訂單檔")]
        errors = []
        for f in xlsx:
            try:
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                ws = wb["Sales"] if "Sales" in wb.sheetnames else wb.active
                row = next(ws.iter_rows(max_row=1, values_only=True), None)
                if not row:
                    errors.append(RowError(0, "sheet", f.name, "工作表為空",
                                           source_file=f.name))
                    continue
                hdrs = {str(v).strip() for v in row if v}
                needed = {"訂單號碼", "商品貨號", "全家服務編號 / 7-11 店號"}
                missing = needed - hdrs
                if missing:
                    errors.append(RowError(0, "header", f.name,
                                           f"缺少欄位：{', '.join(missing)}",
                                           source_file=f.name))
            except Exception as e:
                errors.append(RowError(0, "file", f.name, f"無法開啟：{e}",
                                       source_file=f.name))
        return errors

    def _process(self, input_files: list[Path]) -> ConversionResult:
        # 從 product_map 建立 SKU 直查表（品號 → row_dict）
        sku_map: dict[str, dict] = {}
        for row in self._product_map.values():
            sku = row.get("品號", "")
            if sku and sku not in sku_map:
                sku_map[sku] = row

        xlsx_files    = [f for f in input_files if f.suffix.lower() == ".xlsx"]
        total_success = 0
        total_fail    = 0
        all_errors:   list[RowError] = []
        output_files: list[Path]     = []
        order_date    = ""

        for order_file in xlsx_files:
            wb = openpyxl.load_workbook(order_file, read_only=True, data_only=True)
            ws = wb["Sales"] if "Sales" in wb.sheetnames else wb.active
            all_rows = list(ws.iter_rows(values_only=True))
            if not all_rows:
                continue

            hdr = [str(v).strip() if v else "" for v in all_rows[0]]
            idx = {name: i for i, name in enumerate(hdr)}

            def get(row, col, default=None):
                i = idx.get(col)
                return row[i] if i is not None and i < len(row) else default

            # 從第一筆有值的 發票開立日期 提取訂單日期
            if not order_date:
                for row in all_rows[1:]:
                    d = get(row, "發票開立日期")
                    if d:
                        m = _DATE_SLASH.search(str(d))
                        if m:
                            order_date = m.group(1) + m.group(2) + m.group(3)
                            break

            # 日期 fallback：取檔名前 4 碼作為 MMDD
            if not order_date:
                stem_match = re.match(r'^(\d{4})', order_file.stem)
                if stem_match:
                    order_date = datetime.today().strftime("%Y") + stem_match.group(1)

            output_rows: list[tuple] = []

            for row_idx, row in enumerate(all_rows[1:], start=2):
                if not any(v is not None for v in row):
                    continue

                raw_code = get(row, "商品貨號")
                raw_name = get(row, "商品名稱") or ""

                # 贈品過濾（貨號為空）
                if raw_code is None or str(raw_code).strip() == "":
                    continue

                raw_code = str(raw_code).strip()

                # 品號查找：SKU 直查 → 短代碼正規化 → 品名查找
                # _code_map 每個 key 存 list[dict]（同代碼多規格），取第一筆
                prod = sku_map.get(raw_code)
                if prod is None:
                    code_hits = self._code_map.get(_normalize_code(raw_code))
                    prod = code_hits[0] if code_hits else None
                if prod is None:
                    clean_name = raw_name.replace('團購限定B|', '').strip()
                    prod, candidates = self._lookup_by_name(clean_name)
                    if prod is None:
                        all_errors.append(RowError(
                            row_number=row_idx,
                            field_name="商品貨號",
                            original_value=raw_code,
                            reason=f"找不到品號對照（商品名：{raw_name}）",
                            candidates=candidates,
                            source_file=order_file.name,
                        ))
                        total_fail += 1
                        continue

                # 數量與價格轉換
                try:
                    pack = int(prod.get("包數") or 1) or 1
                except (ValueError, TypeError):
                    pack = 1
                try:
                    qty = int(get(row, "數量") or 1)
                except (ValueError, TypeError):
                    qty = 1
                try:
                    price = float(get(row, "商品結帳價") or 0)
                except (ValueError, TypeError):
                    price = 0.0

                clean_name = (raw_name.replace('團購限定B|', '').strip()
                              if raw_name else prod.get("品名", ""))
                addon = get(row, "加購品類型")
                addon = 0 if addon is None else addon

                def _safe(v, default=""):
                    return v if v is not None else default

                output_rows.append((
                    get(row, "訂單號碼"),
                    get(row, "收件人"),
                    get(row, "完整地址"),
                    get(row, "收件人電話號碼"),
                    get(row, "發票號碼"),
                    prod.get("品號", raw_code),
                    clean_name,
                    qty * pack,
                    0.0 if price == 0 else price / pack,
                    _safe(get(row, "商品折扣優惠"), 0),
                    _safe(get(row, "商品折扣金額"), ""),
                    addon,
                    _safe(get(row, "點數折現分攤"), ""),
                    _safe(get(row, "出貨備註"), ""),
                    _safe(get(row, "全家服務編號 / 7-11 店號"), ""),
                    "",  # 付款方式
                ))
                total_success += 1

            if not output_rows:
                continue

            try:
                date_obj = datetime.strptime(order_date, "%Y%m%d")
            except (ValueError, TypeError):
                date_obj = datetime.today()

            mmdd     = date_obj.strftime("%m%d")
            delivery = _delivery_type(order_file.name)
            out_path = _write_xlsx(self.output_dir.parent, self.source_type,
                                   date_obj, f"{mmdd}{delivery}.xlsx", output_rows)
            output_files.append(out_path)

        if not order_date:
            order_date = datetime.today().strftime("%Y%m%d")

        return ConversionResult(
            source_type   = self.source_type,
            success_count = total_success,
            fail_count    = total_fail,
            order_date    = order_date,
            output_files  = output_files,
            errors        = all_errors,
            status        = "completed" if total_success > 0 else "failed",
        )


def _write_xlsx(base_dir: Path, source_type: str, date_obj: datetime,
                filename: str, rows: list[tuple]) -> Path:
    p = output_path(base_dir, source_type, date_obj, filename)
    wb = Workbook()
    ws = wb.active
    ws.title = date_obj.strftime("%Y-%m-%d")
    ws.append(OUTPUT_HEADER)
    for row in rows:
        ws.append(list(row))
    wb.save(p)
    return p

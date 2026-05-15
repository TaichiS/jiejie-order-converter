"""
converters/jjofficial.py
捷捷寶寶粥官網訂單轉換器。
輸入：後台匯出 xlsx（工作表 Sales）
輸出：MMDD黑貓.xlsx + MMDD全家.xlsx + MMDD黑貓托運單.csv
"""
from __future__ import annotations

import csv
import re
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl import Workbook

from converters.base import BaseConverter, ConversionResult, RowError
from app.utils import output_path

FREIGHT_SKU  = "F59900001"
FREIGHT_NAME = "運費"

SENDER = {
    "name":    "濬詮股份有限公司",
    "tel":     "05-5870993",
    "mobile":  "0909-870-993",
    "address": "雲林縣莿桐鄉大美村溪美22-3號1樓",
    "temp":    "3",
    "size":    "2",
    "deliver": "4",
}

OUTPUT_HEADER = [
    "訂單號碼", "收件人", "完整地址", "收件人電話號碼", "發票號碼",
    "商品貨號", "商品名稱", "數量", "商品結帳價", "商品折扣優惠",
    "商品折扣金額", "點數折現分攤", "出貨備註", "送貨編號", "付款方式",
]

CSV_HEADER = [
    "收件人姓名", "收件人電話", "收件人手機", "收件人地址",
    "代收金額或到付", "件數", "品名(詳參數表)", "備註", "訂單編號",
    "希望配達時間((詳參數表))", "出貨日期(YYYY/MM/DD)", "預定配達日期(YYYY/MM/DD)",
    "溫層((詳參數表))", "尺寸((詳參數表))",
    "寄件人姓名", "寄件人電話", "寄件人手機", "寄件人地址",
    "保值金額(20001~10萬之間)-會產生額外費用", "品名說明", "是否列印",
    "是否捐贈", "統一編號", "手機載具", "愛心碼", "可刷卡(Y/N)", "手機支付(Y/N)",
]

# 非標準品號 → 標準品號對照（來源：官網 CLAUDE.md）
CODE_MAP: dict[str, str] = {
    "G-0-1":          "F50100001",
    "G-0-2":          "F50100002",
    "G-0-4":          "F50100003",
    "G-1-04":         "F50100081",
    "G-1-06":         "F50100005",
    "G-1P-01":        "F50100007",
    "G-1P-05":        "F50100007",
    "G-1P-13":        "F50100009",
    "G-1P-14":        "F50100010",
    "G-S2-01B":       "F50100016",
    "G-S2-02A":       "F50100018",
    "G-S2-02B":       "F50100017",
    "G-S2-09R":       "F50100020",
    "G-S2-10R":       "F50100019",
    "G-S2-5A5B":      "F50100013",
    "GYN09":          "F50100014",
    "N-1":            "F50200002",
    "N-1-3":          "D51300001",
    "1P-05-150":      "D50300005",
    "2-M3":           "D50600012",
    "1-06":           "D50200006",
    "4710586224054":  "F50100022",
    "4710586223002":  "E52010007",
}

# 贈品代號 → 品號對照（商品貨號為空且名稱為「贈AXX*N份」格式時使用）
GIFT_NAME_MAP: dict[str, str] = {
    "A02": "D50800002",
    "A10": "D50800010",
}


class JJOfficialConverter(BaseConverter):

    @property
    def source_type(self) -> str:
        return "jjofficial"

    def _validate(self, input_files: list[Path]) -> list[RowError]:
        xlsx = [f for f in input_files if f.suffix.lower() == ".xlsx"]
        if not xlsx:
            return [RowError(0, "file", "", "找不到 .xlsx 訂單檔")]
        try:
            wb = openpyxl.load_workbook(xlsx[0], read_only=True, data_only=True)
            if "Sales" not in wb.sheetnames:
                return [RowError(0, "sheet", xlsx[0].name, "工作表 'Sales' 不存在",
                                 source_file=xlsx[0].name)]
        except Exception as e:
            return [RowError(0, "file", xlsx[0].name, f"無法開啟：{e}",
                             source_file=xlsx[0].name)]
        return []

    def _process(self, input_files: list[Path]) -> ConversionResult:
        order_file = next(f for f in input_files if f.suffix.lower() == ".xlsx")
        wb_in  = openpyxl.load_workbook(order_file, read_only=True, data_only=True)
        ws_in  = wb_in["Sales"]
        all_rows = [r for r in ws_in.iter_rows(values_only=True)
                    if any(v is not None for v in r)]

        # 動態欄位索引（相容不同匯出版本）
        cm = {str(v).strip(): i for i, v in enumerate(all_rows[0]) if v is not None}

        # 品號索引（品號 → product dict）
        sku_map = {
            v.get("品號", "").strip(): v
            for v in self._product_map.values()
            if v.get("品號", "").strip()
        }

        # 取訂單日期（從第一筆發票開立日期）
        date_col   = cm.get("發票開立日期")
        order_date = ""
        for raw in all_rows[1:]:
            d = raw[date_col] if date_col is not None and len(raw) > date_col else None
            if d:
                order_date = (d.strftime("%Y%m%d") if hasattr(d, "strftime")
                              else re.sub(r"[/-]", "", str(d)[:10]))
                break
        try:
            date_obj = datetime.strptime(order_date, "%Y%m%d")
        except (ValueError, TypeError):
            date_obj   = datetime.today()
            order_date = date_obj.strftime("%Y%m%d")
        mmdd = order_date[4:8]

        # 依訂單號碼收集資料（保持原始順序）
        orders: dict[str, dict] = {}
        errors: list[RowError]  = []
        success_count = 0
        fail_count    = 0

        for row_idx, raw in enumerate(all_rows[1:], start=2):
            def cell(col_name: str, default=None):
                idx = cm.get(col_name)
                return raw[idx] if idx is not None and idx < len(raw) else default

            order_id = str(cell("訂單號碼") or "").strip()
            if not order_id:
                continue

            rcv_name   = str(cell("收件人") or "").strip()
            rcv_phone  = _normalize_phone(str(cell("收件人電話號碼") or ""))
            invoice    = str(cell("發票號碼") or "").strip() or None
            raw_addr   = str(cell("完整地址") or "")
            family_no  = str(cell("全家服務編號 / 7-11 店號") or "").strip()
            is_family  = bool(family_no)
            address    = None if is_family else _parse_address(raw_addr)
            ship_no    = str(cell("送貨編號") or "").strip() or None
            note       = str(cell("出貨備註") or "").strip() or None
            discount   = int(_num(cell("商品折扣金額")))
            points     = int(_num(cell("點數折現分攤")))
            freight    = _num(cell("運費"))
            is_addon   = str(cell("加購品類型") or "").strip() == "主商品加購品"
            raw_sku    = str(cell("商品貨號") or "").strip()
            input_qty  = int(_num(cell("數量") or 1))
            input_price = _num(cell("商品結帳價") or 0)

            # 空品號自動補全（Pattern A / B），補不到才跳過
            is_gift = False
            if not raw_sku:
                item_name = str(cell("商品名稱") or "").strip()
                # Pattern A：名稱開頭就是品號（如 E52990001 母親節滿額贈...）
                m_a = re.match(r"^\s*([A-Z]\d{8})", item_name)
                # Pattern B：贈品（如 贈A02*1份）
                m_b = re.match(r"^\s*贈\s*([A-Z]\d{1,3})\s*\*\s*(\d+)\s*[份包]", item_name)
                if m_a:
                    raw_sku = m_a.group(1)
                elif m_b:
                    gift_code = m_b.group(1)
                    mapped = GIFT_NAME_MAP.get(gift_code)
                    if mapped:
                        raw_sku = mapped
                        is_gift = True
                    else:
                        errors.append(RowError(row_idx, "商品貨號", item_name,
                                               f"贈品代號 {gift_code!r} 在 GIFT_NAME_MAP 中找不到",
                                               source_file=order_file.name))
                        fail_count += 1
                        continue
                else:
                    continue

            # 初始化訂單容器
            if order_id not in orders:
                orders[order_id] = {
                    "rows":      [],
                    "freight":   freight,
                    "is_family": is_family,
                    "rcv_name":  rcv_name,
                    "rcv_phone": rcv_phone,
                    "address":   address,
                }
            elif freight > 0:
                orders[order_id]["freight"] = freight

            # 品號正規化：CODE_MAP → 去除零補位（2-01→2-1）→ 直接查 sku_map
            sku = CODE_MAP.get(raw_sku) or _depad_code(raw_sku, self._code_map) or raw_sku
            product = sku_map.get(sku)

            # 數量倍率：名稱含「N包」或「贈AXX*N份」時乘上 N
            item_name_for_qty = str(cell("商品名稱") or "").strip()
            m_qty = re.search(r"(\d+)\s*[份包]", item_name_for_qty)
            qty_multiplier = int(m_qty.group(1)) if m_qty and not is_addon else 1
            input_qty = input_qty * qty_multiplier

            if is_addon:
                # 加購品：從「寶寶粥官網-加購」查價格、包數與折扣
                addon = self.repository.lookup_by_sku(sku, "寶寶粥官網-加購")
                if not addon:
                    errors.append(RowError(row_idx, "商品貨號", raw_sku,
                                           f"加購品 {sku!r} 在加購品資料中找不到",
                                           source_file=order_file.name))
                    fail_count += 1
                    continue
                pack       = int(_num(addon.get("包數") or 1) or 1)
                output_qty = input_qty * pack
                orig_price = _num(addon.get("份數價格") or 0)
                unit_price = orig_price / pack if pack else orig_price
                addon_disc = int(_num(addon.get("折扣") or 0))
            else:
                addon_disc = 0
                pack = int(_num(product.get("包數") or 0) or 0) if product else 0
                if product:
                    ref_price = _num(product.get("商品結帳價") or 0)
                    if pack > 0:
                        # D/E 系列：展開包數，使用品號資料單包價
                        output_qty = input_qty * pack
                        unit_price = ref_price if ref_price else input_price / pack
                    else:
                        # 組合包（F 系列）：數量不變，價格用品號資料或輸入值
                        output_qty = input_qty
                        unit_price = ref_price if ref_price else input_price
                else:
                    # 品號找不到但 F 系列：直接通過
                    if sku.startswith("F"):
                        output_qty = input_qty
                        unit_price = input_price
                    else:
                        errors.append(RowError(row_idx, "商品貨號", raw_sku,
                                               f"品號 {sku!r} 在品號資料中找不到",
                                               source_file=order_file.name))
                        fail_count += 1
                        continue

            # 贈品（Pattern B）價格強制為 0
            if is_gift:
                unit_price = 0.0

            prod_name = product.get("品名", sku) if product else sku

            orders[order_id]["rows"].append({
                "order_id":  order_id,
                "rcv_name":  rcv_name,
                "address":   address,
                "rcv_phone": rcv_phone,
                "invoice":   invoice,
                "sku":       sku,
                "name":      prod_name,
                "qty":       output_qty,
                "price":     unit_price,
                "discount":  discount,
                "addon_disc": addon_disc,
                "points":    points,
                "note":      note,
                "ship_no":   ship_no,
                "is_family": is_family,
            })
            success_count += 1

        # 拆分黑貓 / 全家，附加運費列，建立托運單資料
        black_rows:  list[list] = []
        family_rows: list[list] = []
        cat_orders:  list[dict] = []

        for order_id, o in orders.items():
            item_rows = list(o["rows"])

            # 同訂單折扣加總（放第一列，其餘為 0）
            total_discount = sum(
                r["discount"] + r["addon_disc"] + r["points"] for r in item_rows
            )

            # 同訂單備註合併（去重後以「；」連接）
            seen_notes: list[str] = []
            for r in item_rows:
                n = r.get("note") or ""
                if n and n not in seen_notes:
                    seen_notes.append(n)
            combined_note = "；".join(seen_notes) or None

            # 運費列附在訂單最後
            if o["freight"] > 0:
                item_rows.append({
                    "order_id":  order_id,
                    "rcv_name":  o["rcv_name"],
                    "address":   o["address"],
                    "rcv_phone": o["rcv_phone"],
                    "invoice":   None,
                    "sku":       FREIGHT_SKU,
                    "name":      FREIGHT_NAME,
                    "qty":       1,
                    "price":     int(o["freight"]),
                    "discount":  0,
                    "addon_disc": 0,
                    "points":    0,
                    "note":      None,
                    "ship_no":   None,
                    "is_family": o["is_family"],
                })

            for i, r in enumerate(item_rows):
                is_first_row = (i == 0)
                row = [
                    r["order_id"], r["rcv_name"], r["address"], r["rcv_phone"],
                    r["invoice"],
                    r["sku"], r["name"], r["qty"], r["price"],
                    0, total_discount if is_first_row else 0, 0,
                    combined_note if r["sku"] != FREIGHT_SKU else None,
                    r["ship_no"], None,
                ]
                if r["is_family"]:
                    family_rows.append(row)
                else:
                    black_rows.append(row)

            if not o["is_family"]:
                first_note = next(
                    (r["note"] for r in o["rows"] if r.get("note")), None
                )
                cat_orders.append({
                    "rcv_name":  o["rcv_name"],
                    "rcv_phone": o["rcv_phone"],
                    "address":   o["address"] or "",
                    "order_id":  order_id,
                    "note":      first_note or "",
                })

        out_black  = _write_xlsx(self.output_dir.parent, self.source_type,
                                 date_obj, f"{mmdd}黑貓.xlsx", black_rows)
        out_family = _write_xlsx(self.output_dir.parent, self.source_type,
                                 date_obj, f"{mmdd}全家.xlsx", family_rows)
        out_csv    = _write_csv(self.output_dir.parent, self.source_type,
                                date_obj, f"{mmdd}黑貓托運單.csv", cat_orders, date_obj)

        status = "completed" if fail_count == 0 else ("partial" if success_count > 0 else "failed")
        return ConversionResult(
            source_type   = self.source_type,
            success_count = success_count,
            fail_count    = fail_count,
            order_date    = order_date,
            output_files  = [out_black, out_family, out_csv],
            errors        = errors,
            status        = status,
        )


# ── 輸出 ──────────────────────────────────────────────────────────────────

def _write_xlsx(base_dir: Path, source_type: str, date_obj: datetime,
                filename: str, rows: list[list]) -> Path:
    p = output_path(base_dir, source_type, date_obj, filename)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(OUTPUT_HEADER)
    for r in rows:
        ws.append(r)
    wb.save(p)
    return p


def _write_csv(base_dir: Path, source_type: str, date_obj: datetime,
               filename: str, orders: list[dict], ref_date: datetime) -> Path:
    p = output_path(base_dir, source_type, date_obj, filename)
    ship_date = ref_date.strftime("%Y/%m/%d")
    next_date = (ref_date + timedelta(days=1)).strftime("%Y/%m/%d")
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(CSV_HEADER)
        for o in orders:
            w.writerow([
                o["rcv_name"], "", o["rcv_phone"], o["address"],
                "", 1, 1, o.get("note") or "1",
                o["order_id"],
                SENDER["deliver"],
                ship_date, next_date,
                SENDER["temp"], SENDER["size"],
                SENDER["name"], SENDER["tel"], SENDER["mobile"], SENDER["address"],
                "", "", "", "", "", "", "", "", "",
            ])
    return p


# ── 工具 ─────────────────────────────────────────────────────────────────

def _parse_address(raw: str) -> str | None:
    """
    解析官網地址格式：'台灣 350 苗栗縣 竹南鎮 中華路...'
    去除「台灣」前綴、郵遞區號，合併城市＋行政區＋街道（無空格）。
    """
    if not raw:
        return None
    s = re.sub(r"^台灣\s*", "", raw.strip())  # 去除「台灣」前綴
    s = re.sub(r"^\d{3,6}\s*", "", s)          # 去除郵遞區號
    s = s.replace(" ", "")                     # 去除空格
    return s if s else None


def _normalize_phone(raw: str) -> str:
    """正規化電話：去除 +886 國碼，補回 0 前綴，取前 10 碼。"""
    s = re.sub(r"\s+", "", raw.strip())
    if s.startswith("+886"):
        s = "0" + s[4:]
    elif s.startswith("886") and len(s) > 9:
        s = "0" + s[3:]
    return s[:10] if len(s) >= 10 else s


def _depad_code(raw: str, code_map: dict) -> str | None:
    """
    在 code_map 中查詢品號，回傳對應品號或 None。
    先直查（如 1-01 → D50200001），找不到再去除零補位（如 2-01 → 2-1）。
    """
    product = code_map.get(raw)
    if product:
        return product.get("品號", "").strip() or None
    normalized = re.sub(r"-0+(\d)", r"-\1", raw)
    if normalized != raw:
        product = code_map.get(normalized)
        if product:
            return product.get("品號", "").strip() or None
    return None


def _num(val) -> float:
    try:
        return float(str(val).strip())
    except (ValueError, TypeError):
        return 0.0

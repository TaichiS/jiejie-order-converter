"""
converters/kadomo.py
卡多摩嬰童館採購單轉換器。
輸入：YYYYMMDD_採購單[倉別] [店名].xlsx（含「倉別」標題列）
輸出：{倉別}-{MMDD}.xlsx（新細明體 12pt，固定欄寬）
參考資料：卡多摩/通路資料.csv（編號精確查找）
"""
from __future__ import annotations

import csv
import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font

from converters.base import BaseConverter, ConversionResult, RowError
from app.utils import output_path

HEADERS = ["訂單編號", "收件人", "地址", "電話", "產品編號", "產品名稱", "數量", "單價", "備註", "備註.1"]

COL_WIDTHS = {
    "A": 13.875, "B": 15.0,  "C": 26.125,
    "D": 14.25,  "E": 11.0,  "F": 28.125,
    "G": 6.0,    "H": 8.0,   "I": 18.0,
}

OUTPUT_FONT = Font(name="新細明體", size=12)

_DATE_RE = re.compile(r"(\d{8})")  # YYYYMMDD from filename


class KadomoConverter(BaseConverter):

    def __init__(self, repository, output_dir: Path):
        super().__init__(repository, output_dir)
        self.data_dir = Path(__file__).resolve().parent.parent / "卡多摩"
        self._barcode_map: dict[str, str]   = {}  # barcode → product_code
        self._price_map:   dict[str, float] = {}  # product_code → price
        self._store_map:   dict[str, dict]  = {}  # 編號 → {店名, 電話, 地址}

    @property
    def source_type(self) -> str:
        return "kadomo"

    def _load_reference(self) -> None:
        # 條碼與定價從 Repository（已匯入 DB）
        self._barcode_map = self.repository.load_barcodes("kadomo")
        from app.models import ChannelPrice
        self._price_map = {
            cp.sku: float(cp.price)
            for cp in ChannelPrice.query.filter_by(channel="kadomo").all()
        }
        # 通路資料：以倉別編號精確查找
        csv_path = self.data_dir / "通路資料.csv"
        with open(csv_path, newline="", encoding="utf-8-sig") as f:
            self._store_map = {
                row["編號"]: row
                for row in csv.DictReader(f)
                if row["編號"]
            }

    # ── 驗證 ──────────────────────────────────────────────────────────────

    def _validate(self, input_files: list[Path]) -> list[RowError]:
        xlsx = [f for f in input_files if f.suffix.lower() == ".xlsx"]
        if not xlsx:
            return [RowError(0, "file", "", "找不到 .xlsx 訂單檔")]
        for f in xlsx:
            try:
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                if _find_header(wb) is None:
                    return [RowError(0, "sheet", f.name,
                                     "找不到含「倉別」與「商品條碼」的標題列",
                                     source_file=f.name)]
            except Exception as e:
                return [RowError(0, "file", f.name, f"無法開啟：{e}",
                                 source_file=f.name)]
        return []

    # ── 轉換 ──────────────────────────────────────────────────────────────

    def _process(self, input_files: list[Path]) -> ConversionResult:
        xlsx_files = [f for f in input_files if f.suffix.lower() == ".xlsx"]

        all_output: list[Path]     = []
        all_errors: list[RowError] = []
        total_success = 0
        total_fail    = 0
        order_date    = ""

        for f in xlsx_files:
            out, success, fail, errors, date = self._process_one(f)
            if out:
                all_output.append(out)
            total_success += success
            total_fail    += fail
            all_errors.extend(errors)
            if date and not order_date:
                order_date = date

        status = "completed" if total_fail == 0 else ("partial" if total_success > 0 else "failed")

        return ConversionResult(
            source_type   = self.source_type,
            success_count = total_success,
            fail_count    = total_fail,
            order_date    = order_date,
            output_files  = all_output,
            errors        = all_errors,
            status        = status,
        )

    def _process_one(self, f: Path) -> tuple[Path | None, int, int, list[RowError], str]:
        wb       = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws       = wb.active
        all_rows = list(ws.iter_rows(values_only=True))

        header_idx, col_map = _find_header(wb)
        data_rows = all_rows[header_idx + 1:]

        # 日期：優先從檔名提取
        order_date = _extract_date(f.name)

        # 倉別代碼：取第一個非空資料列
        w_col = col_map.get("倉別")
        warehouse_code = ""
        for row in data_rows:
            v = _cell(row, w_col)
            if v:
                warehouse_code = str(v).strip()
                break

        mmdd       = order_date[4:8] if len(order_date) >= 8 else "0000"
        order_id   = f"{warehouse_code}-{mmdd}" if warehouse_code else mmdd

        # 門市資訊
        store_info, store_error = self._resolve_store_info(warehouse_code, f.name)
        if store_error:
            return None, 0, 1, [store_error], order_date

        recipient = store_info["店名"]
        address   = store_info["地址"]
        phone     = store_info["電話"]

        # 欄位索引
        b_col      = col_map.get("商品條碼")
        n_col      = col_map.get("商品名稱")
        q_col      = col_map.get("進貨數量")
        po_col     = col_map.get("採購單號")
        remark_col = col_map.get("商品備註")

        output_rows: list[tuple] = []
        errors:      list[RowError] = []

        for row_idx, row in enumerate(data_rows, start=header_idx + 2):
            barcode = _cell_str(row, b_col)
            name    = _cell_str(row, n_col)
            qty     = _cell(row, q_col)
            po_no   = _cell_str(row, po_col)
            remark  = _cell_str(row, remark_col)

            # 跳過空列
            if not barcode and not name:
                continue

            # 條碼查品號
            product_code = self._barcode_map.get(barcode, "") if barcode else ""
            if barcode and not product_code:
                errors.append(RowError(
                    row_idx, "商品條碼", barcode,
                    "條碼對照表中找不到對應品號",
                    source_file=f.name,
                ))

            price = self._price_map.get(product_code, 0.0) if product_code else 0.0

            output_rows.append((
                order_id, recipient, address, phone,
                product_code, name, qty, price, po_no, remark,
            ))

        if not output_rows:
            return None, 0, len(errors), errors, order_date

        try:
            date_obj = datetime.strptime(order_date, "%Y%m%d")
        except ValueError:
            date_obj = datetime.today()
            order_date = date_obj.strftime("%Y%m%d")

        out_path = _write_xlsx(
            self.output_dir.parent, self.source_type,
            date_obj, f"{order_id}.xlsx", output_rows,
        )
        return out_path, len(output_rows), len(errors), errors, order_date

    def _resolve_store_info(self, warehouse_code: str, filename: str) -> tuple[dict | None, RowError | None]:
        store_info = self._store_map.get(warehouse_code)
        if not store_info:
            return None, RowError(
                0, "倉別", warehouse_code,
                f"倉別編號 {warehouse_code!r} 不在通路資料.csv 中",
                source_file=filename,
            )
        return store_info, None


# ── 輸出 ──────────────────────────────────────────────────────────────────

def _write_xlsx(base_dir: Path, source_type: str, date_obj: datetime,
                filename: str, rows: list[tuple]) -> Path:
    p  = output_path(base_dir, source_type, date_obj, filename)
    wb = Workbook()
    ws = wb.active
    ws.title = date_obj.strftime("%Y-%m-%d")

    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = OUTPUT_FONT

    for row in rows:
        ws.append(list(row))
        for cell in ws[ws.max_row]:
            cell.font = OUTPUT_FONT

    wb.save(p)
    return p


# ── 工具 ──────────────────────────────────────────────────────────────────

def _find_header(wb: openpyxl.Workbook) -> tuple[int, dict[str, int]] | None:
    """
    搜尋前 15 列，找到同時含有「倉別」與「商品條碼」的列。
    回傳 (列索引, {欄名: 欄位index}) 或 None。
    """
    ws = wb.active
    rows = list(ws.iter_rows(max_row=15, values_only=True))
    for i, row in enumerate(rows):
        cells = [str(v).strip() if v is not None else "" for v in row]
        if "倉別" in cells and "商品條碼" in cells:
            return i, {v: j for j, v in enumerate(cells) if v}
    return None


def _extract_date(filename: str) -> str:
    """從檔名提取 YYYYMMDD，失敗則回傳今日"""
    m = _DATE_RE.search(filename)
    if m:
        candidate = m.group(1)
        try:
            datetime.strptime(candidate, "%Y%m%d")
            return candidate
        except ValueError:
            pass
    return datetime.today().strftime("%Y%m%d")


def _cell(row: tuple, idx: int | None):
    """安全取值"""
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _cell_str(row: tuple, idx: int | None) -> str:
    v = _cell(row, idx)
    return str(v).strip() if v is not None else ""

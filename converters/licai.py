"""
converters/licai.py
麗兒采家嬰童館採購單轉換器。

支援兩種輸入格式：
  舊格式（split_orders.py 拆分後）：每個 xlsx 為單一門市
    Row 0: ('門市：短名', None, '門市資訊：地址 / 電話', ...)
    Row 1: ('單號：UWXXXXXX', None, '採購日期：YYYY/M/D', ...)
    Row 2: 空列
    Row 3+: (序號, 條碼, 品名, 品牌, 市價, 數量)

  新格式（合併採購單）：一個 xlsx 包含多家門市，每段重複：
    Row N+0: ('麗兒采家採購單 - 供應商', ...)  ← 標題列（可選）
    Row N+1: ('門市：短名', None, '門市資訊：...', ...)
    Row N+2: ('單號：UWXXXXXX', None, '採購日期：YYYY/M/D', ...)
    Row N+3: (序, 貨號, 品名, 品牌, 市價, 門市短名)  ← 欄位標題列
    Row N+4+: (序號, 條碼, 品名, 品牌, 市價, 數量)

輸出：MMDD 麗采{門市短名}.xlsx（每家門市一個檔案）
  欄位：訂單編號, 收件人, 地址, 電話, 產品編號, 產品名稱, 數量, 單價, 備註, 備註.1

參考資料：麗兒采家/通路資料.csv（店名精確查找）
"""
from __future__ import annotations

import csv
import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font

from converters.base import BaseConverter, ConversionResult, RowError
from app.utils import output_path

HEADERS = ["訂單編號", "收件人", "地址", "電話", "產品編號", "產品名稱", "數量", "單價", "備註", "備註.1"]
OUTPUT_FONT = Font(name="新細明體", size=12)

_STORE_RE = re.compile(r"門市：(\S+)")
_ORDER_RE = re.compile(r"單號：(\S+)")
_DATE_RE  = re.compile(r"採購日期：(\d{4})/(\d{1,2})/(\d{1,2})")
_STORE_SHORT_RE = re.compile(r"麗兒采家-(.+?)店")


class LicaiConverter(BaseConverter):

    def __init__(self, repository, output_dir: Path):
        super().__init__(repository, output_dir)
        self.data_dir = Path(__file__).resolve().parent.parent / "麗兒采家"
        self._barcode_map: dict[str, str]   = {}  # barcode → sku
        self._name_map:    dict[str, str]   = {}  # sku → 品名
        self._price_map:   dict[str, float] = {}  # sku → 包數價格
        self._pack_map:    dict[str, int]   = {}  # sku → 包數
        self._store_map:   dict[str, dict]  = {}  # 短名 → {store_full, phone, address}

    @property
    def source_type(self) -> str:
        return "licai"

    def _load_reference(self) -> None:
        self._barcode_map = self.repository.load_barcodes("licai")
        from app.models import Product, ChannelPrice, UnifiedProduct
        sku_set = {cp.sku for cp in ChannelPrice.query.filter_by(channel="licai").all()}
        products = Product.query.filter(Product.sku.in_(sku_set)).all()
        self._name_map = {p.sku: p.name for p in products}
        # 從 unified_products 麗采通路取得最新包數價格與包數
        licai_products = UnifiedProduct.query.filter_by(channel="麗采").all()
        self._price_map = {
            up.sku: float(up.pack_price)
            for up in licai_products
            if up.pack_price is not None
        }
        self._pack_map = {
            up.sku: int(up.pack_size)
            for up in licai_products
            if up.pack_size is not None and up.pack_size > 0
        }
        self._store_map = self._load_store_map()

    def _load_store_map(self) -> dict[str, dict]:
        csv_path = self.data_dir / "通路資料.csv"
        store_map: dict[str, dict] = {}
        with open(csv_path, newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                store_full = row["店名"]
                m = _STORE_SHORT_RE.search(store_full)
                if not m:
                    continue
                short = m.group(1).replace("旗艦", "")
                store_map[short] = {
                    "store_full": store_full,
                    "phone":      row["電話"],
                    "address":    row["地址"],
                }
        return store_map

    def _validate(self, input_files: list[Path]) -> list[RowError]:
        xlsx = [f for f in input_files if f.suffix.lower() == ".xlsx"]
        if not xlsx:
            return [RowError(0, "file", "", "找不到 .xlsx 採購單")]
        for f in xlsx:
            try:
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(max_row=5, values_only=True))
                if not rows or not any(
                    any(v and "門市：" in str(v) for v in row) for row in rows
                ):
                    return [RowError(0, "sheet", f.name,
                                     "找不到「門市：」標頭列",
                                     source_file=f.name)]
            except Exception as e:
                return [RowError(0, "file", f.name, f"無法開啟：{e}",
                                 source_file=f.name)]
        return []

    def _process(self, input_files: list[Path]) -> ConversionResult:
        xlsx_files = [f for f in input_files if f.suffix.lower() == ".xlsx"]
        all_output:  list[Path]     = []
        all_errors:  list[RowError] = []
        total_success = 0
        total_fail    = 0
        order_date    = ""

        for f in xlsx_files:
            outputs, success, fail, errors, date = self._process_one(f)
            all_output.extend(outputs)
            total_success += success
            total_fail    += fail
            all_errors.extend(errors)
            if date and not order_date:
                order_date = date

        return ConversionResult(
            source_type   = self.source_type,
            success_count = total_success,
            fail_count    = total_fail,
            order_date    = order_date,
            output_files  = all_output,
            errors        = all_errors,
            status        = "completed" if total_success > 0 else "failed",
        )

    def _process_one(self, f: Path) -> tuple[list[Path], int, int, list[RowError], str]:
        wb   = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws   = wb.active
        rows = list(ws.iter_rows(values_only=True))

        # 找出所有「門市：」列的索引，每個索引代表一家門市區塊的起點
        block_starts = [
            i for i, row in enumerate(rows)
            if any(v and "門市：" in str(v) for v in row)
        ]

        all_outputs: list[Path]     = []
        all_errors:  list[RowError] = []
        total_success = 0
        total_fail    = 0
        order_date    = ""

        for b, start in enumerate(block_starts):
            end   = block_starts[b + 1] if b + 1 < len(block_starts) else len(rows)
            block = rows[start:end]
            out, success, fail, errors, date = self._process_block(f, block, start)
            if out:
                all_outputs.append(out)
            total_success += success
            total_fail    += fail
            all_errors.extend(errors)
            if date and not order_date:
                order_date = date

        return all_outputs, total_success, total_fail, all_errors, order_date

    def _process_block(
        self, f: Path, rows: list, row_offset: int
    ) -> tuple[Path | None, int, int, list[RowError], str]:
        # rows[0] = 門市：短名 列，rows[1] = 單號：/ 採購日期：列
        store_row_str = " ".join(str(v) for v in rows[0] if v)
        order_row_str = " ".join(str(v) for v in rows[1] if v) if len(rows) > 1 else ""

        store_m = _STORE_RE.search(store_row_str)
        order_m = _ORDER_RE.search(order_row_str)
        date_m  = _DATE_RE.search(order_row_str)

        store_short = store_m.group(1) if store_m else ""
        order_id    = order_m.group(1) if order_m else ""
        order_date  = ""
        if date_m:
            y, mo, d = date_m.groups()
            order_date = f"{y}{int(mo):02d}{int(d):02d}"

        info        = self._store_map.get(store_short, {})
        store_full  = info.get("store_full", f"麗兒采家-{store_short}店")
        phone       = info.get("phone", "")
        address     = info.get("address", "")

        wb_out = Workbook()
        ws_out = wb_out.active
        ws_out.append(HEADERS)

        success = 0
        fail    = 0
        errors: list[RowError] = []

        for i, row in enumerate(rows[2:], start=row_offset + 3):
            if not isinstance(row[0], (int, float)):  # 略過空列與欄位標題列
                continue
            try:
                barcode   = str(row[1]).strip() if row[1] is not None else ""
                xlsx_price = float(row[4]) if row[4] is not None else 0.0
                qty       = int(row[5])   if row[5] is not None else 0
            except (TypeError, ValueError) as e:
                errors.append(RowError(i, "data", str(row), str(e),
                                       source_file=f.name))
                continue

            sku   = self._barcode_map.get(barcode, "")
            name  = self._name_map.get(sku, str(row[2]) if row[2] else "")
            # 優先使用資料庫價格，找不到則 fallback 至 xlsx 市價
            price = self._price_map.get(sku, xlsx_price) if sku else xlsx_price
            # 數量 × 包數（DB 查無包數則 ×1）
            pack  = self._pack_map.get(sku, 1) if sku else 1
            qty   = qty * pack

            if not sku:
                errors.append(RowError(
                    i, "barcode", barcode,
                    f"條碼 {barcode!r} 查無品號",
                    source_file=f.name,
                ))
                fail += 1
            else:
                success += 1

            # 採購日期格式：MMDD
            date_str = order_date[4:8] if len(order_date) == 8 else ""
            ws_out.append([
                order_id, store_full, address, phone,
                sku, name, qty, price, date_str, None,
            ])

        for row in ws_out.iter_rows():
            for cell in row:
                cell.font = OUTPUT_FONT

        mmdd     = order_date[4:8] if len(order_date) == 8 else "0000"
        date_obj = datetime.strptime(order_date, "%Y%m%d") if order_date else datetime.today()
        out_name = f"{mmdd} 麗采{store_short}.xlsx"
        out_path = output_path(self.output_dir.parent, self.source_type, date_obj, out_name)

        wb_out.save(out_path)
        return out_path, success, fail, errors, order_date

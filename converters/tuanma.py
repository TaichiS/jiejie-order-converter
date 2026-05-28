"""
converters/tuanma.py
其他團媽訂單轉換器。
輸入：後台匯出 xlsx（工作表 Sales）
輸出：MMDD其他團媽黑貓.xlsx + MMDD其他團媽全家.xlsx（各 16 欄）

轉換邏輯：
- Sales 工作表；依「全家服務編號 / 7-11 店號」有無分拆黑貓／全家
- 地址：去除「台灣」前綴與郵遞區號，去除空格
- 電話：去除 +886 國碼，補回 0 前綴，截至 10 碼
- 品號查找：SKU 直查 → 品名查找
- 數量 × 包數；價格 ÷ 包數
- 運費 > 0 → 補 F59900001 運費列；附加費 > 0 → 補 F59900003 附加費列
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook

from converters.base import BaseConverter, ConversionResult, RowError
from app.utils import output_path

OUTPUT_HEADER = [
    "訂單號碼", "收件人", "完整地址", "收件人電話號碼", "發票號碼",
    "商品貨號", "商品名稱", "數量", "商品結帳價", "商品折扣優惠",
    "商品折扣金額", "加購折扣", "點數折現分攤", "出貨備註", "送貨編號",
    "付款方式",
]

_DATE_SLASH = re.compile(r'(\d{4})/(\d{2})/(\d{2})')
_TW_PREFIX  = re.compile(r'^台灣\s*')
_POSTAL     = re.compile(r'^\d{3,6}\s+')


def _normalize_address(addr) -> str:
    s = str(addr).strip() if addr else ""
    s = _TW_PREFIX.sub("", s)
    s = _POSTAL.sub("", s)
    return s.replace(" ", "")


def _normalize_phone(phone) -> str:
    s = str(phone).strip() if phone else ""
    if s.startswith("+886"):
        s = "0" + s[4:]
    elif s.startswith("886") and len(s) >= 11:
        s = "0" + s[3:]
    return s[:10]


class TuanmaConverter(BaseConverter):

    @property
    def source_type(self) -> str:
        return "tuanma"

    def _validate(self, input_files: list[Path]) -> list[RowError]:
        xlsx = [f for f in input_files if f.suffix.lower() in (".xlsx", ".xls")]
        if not xlsx:
            return [RowError(0, "file", "", "找不到 .xlsx/.xls 訂單檔")]
        errors = []
        for f in xlsx:
            try:
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                if "Sales" not in wb.sheetnames:
                    errors.append(RowError(0, "sheet", f.name, "工作表 'Sales' 不存在",
                                           source_file=f.name))
                    continue
                ws = wb["Sales"]
                row = next(ws.iter_rows(max_row=1, values_only=True), None)
                if not row:
                    errors.append(RowError(0, "sheet", f.name, "工作表為空",
                                           source_file=f.name))
                    continue
                hdrs = {str(v).strip() for v in row if v}
                needed = {"訂單號碼", "商品貨號", "全家服務編號 / 7-11 店號"}
                missing = needed - hdrs
                if missing:
                    errors.append(RowError(0, "header", f.name,
                                           f"缺少欄位：{', '.join(missing)}",
                                           source_file=f.name))
            except Exception as e:
                errors.append(RowError(0, "file", f.name, f"無法開啟：{e}",
                                       source_file=f.name))
        return errors

    def _process(self, input_files: list[Path]) -> ConversionResult:
        # 建立 SKU 直查表（品號 → row_dict）
        sku_map: dict[str, dict] = {}
        for row in self._product_map.values():
            sku = row.get("品號", "")
            if sku and sku not in sku_map:
                sku_map[sku] = row

        xlsx_files    = [f for f in input_files if f.suffix.lower() in (".xlsx", ".xls")]
        total_success = 0
        total_fail    = 0
        all_errors:   list[RowError] = []
        output_files: list[Path]     = []
        order_date    = ""

        for order_file in xlsx_files:
            wb = openpyxl.load_workbook(order_file, read_only=True, data_only=True)
            ws = wb["Sales"]
            all_rows = list(ws.iter_rows(values_only=True))
            if not all_rows:
                continue

            hdr = [str(v).strip() if v else "" for v in all_rows[0]]
            idx = {name: i for i, name in enumerate(hdr)}

            def get(row, col, default=None):
                i = idx.get(col)
                return row[i] if i is not None and i < len(row) else default

            # 從發票開立日期取訂單日期
            if not order_date:
                for row in all_rows[1:]:
                    d = get(row, "發票開立日期")
                    if d:
                        m = _DATE_SLASH.search(str(d))
                        if m:
                            order_date = m.group(1) + m.group(2) + m.group(3)
                            break

            # fallback：取檔名前 4 碼作為 MMDD
            if not order_date:
                stem_match = re.match(r'^(\d{4})', order_file.stem)
                if stem_match:
                    order_date = datetime.today().strftime("%Y") + stem_match.group(1)

            # 按配送方式分組（品項列 + 暫存各訂單的運費/附加費）
            heikao_rows: list[tuple] = []
            jj_rows:     list[tuple] = []
            # order_id → (freight, extra, base_info_tuple) 用於最後補列
            order_extras: dict[str, dict] = {}

            for row_idx, row in enumerate(all_rows[1:], start=2):
                if not any(v is not None for v in row):
                    continue

                raw_code = get(row, "商品貨號")
                if raw_code is None or str(raw_code).strip() == "":
                    continue
                raw_code = str(raw_code).strip()

                raw_name = get(row, "商品名稱") or ""

                # 品號查找
                prod = sku_map.get(raw_code)
                if prod is None:
                    # 去除前綴：「團購限定A|」、「團購限定B|」或純「團購|」
                    clean_name = re.sub(r'^(?:團購限定[AB]|團購)\|', '', raw_name).strip()
                    # 將剩餘的 | 替換為 -（DB 中以 - 作分隔，如 寶寶義式冰淇淋-藍莓優格）
                    clean_name = clean_name.replace('|', '-')
                    prod, candidates = self._lookup_by_name(clean_name)
                    if prod is None:
                        all_errors.append(RowError(
                            row_number=row_idx,
                            field_name="商品貨號",
                            original_value=raw_code,
                            reason=f"找不到品號對照（商品名：{raw_name}）",
                            candidates=candidates,
                            source_file=order_file.name,
                        ))
                        total_fail += 1
                        continue

                # 數量與價格
                try:
                    pack = int(prod.get("包數") or 1) or 1
                except (ValueError, TypeError):
                    pack = 1
                try:
                    qty = int(get(row, "數量") or 1)
                except (ValueError, TypeError):
                    qty = 1
                try:
                    price = float(get(row, "商品結帳價") or 0)
                except (ValueError, TypeError):
                    price = 0.0

                prod_name = prod.get("品名") or re.sub(r'^(?:團購限定[AB]|團購)\|', '', raw_name).strip().replace('|', '-')
                fjsh = get(row, "全家服務編號 / 7-11 店號")
                fjsh_val = str(fjsh).strip() if fjsh else ""
                order_id  = str(get(row, "訂單號碼") or "")

                out_row = (
                    get(row, "訂單號碼"),
                    get(row, "收件人"),
                    _normalize_address(get(row, "完整地址")),
                    _normalize_phone(get(row, "收件人電話號碼")),
                    get(row, "發票號碼"),
                    prod.get("品號", raw_code),
                    prod_name,
                    qty * pack,
                    0.0 if price == 0 else round(price / pack, 2),
                    0,   # 商品折扣優惠
                    0,   # 商品折扣金額
                    0,   # 加購折扣
                    0,   # 點數折現分攤
                    get(row, "出貨備註") or "",
                    fjsh_val,  # 送貨編號（全家超商用）
                    "",  # 付款方式
                )

                if fjsh_val:
                    jj_rows.append(out_row)
                else:
                    heikao_rows.append(out_row)
                total_success += 1

                # 記錄運費 / 附加費（只記第一次出現的有值列）
                if order_id not in order_extras:
                    try:
                        freight = float(get(row, "運費") or 0)
                    except (ValueError, TypeError):
                        freight = 0.0
                    try:
                        extra = float(get(row, "附加費") or 0)
                    except (ValueError, TypeError):
                        extra = 0.0
                    if freight > 0 or extra > 0:
                        order_extras[order_id] = {
                            "freight": freight,
                            "extra":   extra,
                            "base": (
                                get(row, "訂單號碼"),
                                get(row, "收件人"),
                                _normalize_address(get(row, "完整地址")),
                                _normalize_phone(get(row, "收件人電話號碼")),
                                get(row, "發票號碼"),
                                get(row, "出貨備註") or "",
                                fjsh_val,
                            ),
                        }

            # 在各配送組最後補運費 / 附加費列
            for info in order_extras.values():
                oid, rcv, addr, phone, inv, note, fjsh_val = info["base"]
                if info["freight"] > 0:
                    fr = (oid, rcv, addr, phone, inv,
                          "F59900001", "運費", 1, info["freight"],
                          0, 0, 0, 0, note, fjsh_val, "")
                    (jj_rows if fjsh_val else heikao_rows).append(fr)
                if info["extra"] > 0:
                    ex = (oid, rcv, addr, phone, inv,
                          "F59900003", "附加費", 1, info["extra"],
                          0, 0, 0, 0, note, fjsh_val, "")
                    (jj_rows if fjsh_val else heikao_rows).append(ex)

            try:
                date_obj = datetime.strptime(order_date, "%Y%m%d")
            except (ValueError, TypeError):
                date_obj = datetime.today()

            mmdd = date_obj.strftime("%m%d")

            if heikao_rows:
                p = _write_xlsx(self.output_dir.parent, self.source_type,
                                date_obj, f"{mmdd}其他團媽黑貓.xlsx", heikao_rows)
                output_files.append(p)

            if jj_rows:
                p = _write_xlsx(self.output_dir.parent, self.source_type,
                                date_obj, f"{mmdd}其他團媽全家.xlsx", jj_rows)
                output_files.append(p)

        if not order_date:
            order_date = datetime.today().strftime("%Y%m%d")

        return ConversionResult(
            source_type   = self.source_type,
            success_count = total_success,
            fail_count    = total_fail,
            order_date    = order_date,
            output_files  = output_files,
            errors        = all_errors,
            status        = "completed" if total_success > 0 else "failed",
        )


def _write_xlsx(base_dir: Path, source_type: str, date_obj: datetime,
                filename: str, rows: list[tuple]) -> Path:
    p = output_path(base_dir, source_type, date_obj, filename)
    wb = Workbook()
    ws = wb.active
    ws.title = date_obj.strftime("%Y-%m-%d")
    ws.append(OUTPUT_HEADER)
    for row in rows:
        ws.append(list(row))
    wb.save(p)
    return p

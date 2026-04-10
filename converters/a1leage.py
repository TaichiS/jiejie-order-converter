"""
converters/a1leage.py
樂齡官網訂單轉換器。
輸入：後台匯出 xlsx（工作表 Orders）
輸出：MMDD.xlsx（Sheet1, 15欄）+ MMDD黑貓.csv
"""
from __future__ import annotations

import csv
import re
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl import Workbook

from converters.base import BaseConverter, ConversionResult, RowError
from app.utils import output_path

COD_FEE_SKU   = "F59900003"
COD_FEE_PRICE = 30

SENDER = {
    "name":    "濬詮股份有限公司",
    "tel":     "05-5870993",
    "mobile":  "0909-870-993",
    "address": "雲林縣莿桐鄉大美村溪美22-3號1樓",
    "temp":    "3",
    "size":    "2",
    "deliver": "4",
}

OUTPUT_HEADER = [
    "訂單號碼", "收件人", "完整地址", "收件人電話號碼", "發票號碼",
    "商品貨號", "商品名稱", "數量", "商品結帳價", "商品折扣優惠",
    "商品折扣金額", "點數折現分攤", "出貨備註", "送貨編號", "付款方式",
]

CSV_HEADER = [
    "收件人姓名", "收件人電話", "收件人手機", "收件人地址",
    "代收金額或到付", "件數", "品名(詳參數表)", "備註", "訂單編號",
    "希望配達時間((詳參數表))", "出貨日期(YYYY/MM/DD)", "預定配達日期(YYYY/MM/DD)",
    "溫層((詳參數表))", "尺寸((詳參數表))",
    "寄件人姓名", "寄件人電話", "寄件人手機", "寄件人地址",
    "保值金額(20001~10萬之間)-會產生額外費用", "品名說明", "是否列印",
    "是否捐贈", "統一編號", "手機載具", "愛心碼", "可刷卡(Y/N)", "手機支付(Y/N)",
]


class A1LeageConverter(BaseConverter):

    @property
    def source_type(self) -> str:
        return "a1leage"

    def _validate(self, input_files: list[Path]) -> list[RowError]:
        xlsx = [f for f in input_files if f.suffix.lower() == ".xlsx"]
        if not xlsx:
            return [RowError(0, "file", "", "找不到 .xlsx 訂單檔")]
        try:
            wb = openpyxl.load_workbook(xlsx[0], read_only=True, data_only=True)
            if "Orders" not in wb.sheetnames:
                return [RowError(0, "sheet", xlsx[0].name, "工作表 'Orders' 不存在",
                                 source_file=xlsx[0].name)]
        except Exception as e:
            return [RowError(0, "file", xlsx[0].name, f"無法開啟：{e}",
                             source_file=xlsx[0].name)]
        return []

    def _process(self, input_files: list[Path]) -> ConversionResult:
        order_file = next(f for f in input_files if f.suffix.lower() == ".xlsx")
        wb_in  = openpyxl.load_workbook(order_file, read_only=True, data_only=True)
        ws_in  = wb_in["Orders"]
        all_rows = list(ws_in.iter_rows(values_only=True))

        # 由標題列動態取得欄位索引（應對不同版本匯出格式）
        cm = {str(v).strip(): i for i, v in enumerate(all_rows[0]) if v}

        # 取訂單日期
        order_date = ""
        for raw in all_rows[1:]:
            d = raw[cm.get("訂購日期", 1)]
            if d:
                order_date = d.strftime("%Y%m%d") if hasattr(d, "strftime") else str(d)[:10].replace("-", "")
                break
        try:
            date_obj = datetime.strptime(order_date, "%Y%m%d")
        except (ValueError, TypeError):
            date_obj = datetime.today()
            order_date = date_obj.strftime("%Y%m%d")
        mmdd = order_date[4:8]

        # 品號索引（stripped 品號 → product）
        sku_map = {
            v.get("品號", "").strip(): v
            for v in self._product_map.values()
            if v.get("品號", "").strip()
        }

        xlsx_rows: list[dict]  = []
        csv_orders: list[dict] = []
        errors: list[RowError] = []
        success_count = 0
        fail_count    = 0

        for row_idx, raw in enumerate(all_rows[1:], start=2):
            row = list(raw)

            def cell(col_name: str, default=""):
                idx = cm.get(col_name)
                return row[idx] if idx is not None and idx < len(row) else default

            order_id  = str(cell("訂單編號") or "").strip()
            rcv_name  = str(cell("收件人姓名") or "").strip()
            rcv_phone = str(cell("收件人手機") or "").strip()
            payment   = str(cell("付費方式") or "").strip()
            address   = _parse_address(str(cell("收件人地址") or ""))

            skus       = _split(cell("貨號"))
            item_names = _split(cell("購買品項"))
            qtys       = _split(cell("數量"))
            discount_total = sum(_num(d) for d in _split(cell("折扣金額")))

            # 判斷哪個欄位代表「是否為組合包子品項（應跳過）」：
            #   有 商品單價 欄 → 用商品單價=0 判斷（2月/3月新格式）
            #   否則          → 用商品原價=0 判斷（舊格式）
            if "商品單價" in cm:
                skip_prices = _split(cell("商品單價"))
            else:
                skip_prices = _split(cell("商品原價"))

            is_first   = True
            has_error  = False
            order_subtotal = 0.0

            for i, sku in enumerate(skus):
                skip_price = _num(skip_prices[i] if i < len(skip_prices) else 0)
                if skip_price == 0:
                    continue  # 跳過組合包子品項

                input_qty  = int(_num(qtys[i])) if i < len(qtys) else 1
                raw_name   = item_names[i] if i < len(item_names) else ""
                clean_name = re.sub(r"^捷捷樂齡食品-\s*", "", raw_name).strip()

                product = sku_map.get(sku.strip()) or self._product_map.get(clean_name)
                if not product:
                    errors.append(RowError(
                        row_idx, "貨號", sku,
                        f"品號 {sku!r}（{clean_name}）在品號資料中找不到",
                        source_file=order_file.name,
                    ))
                    has_error = True
                    fail_count += 1
                    continue

                prod_no    = product.get("品號", sku).strip()
                prod_name  = product.get("品名", clean_name)
                pack_count = int(_num(product.get("包數") or 1) or 1)
                unit_price = (_num(product.get("份數價格")) if product.get("份數價格")
                              else _num(product.get("商品結帳價") or 0))
                output_qty = input_qty * pack_count

                xlsx_rows.append({
                    "order_id": order_id,
                    "rcv_name": rcv_name,
                    "address":  address,
                    "rcv_phone": rcv_phone,
                    "sku":      prod_no,
                    "name":     prod_name,
                    "qty":      output_qty,
                    "price":    unit_price,
                    "discount": int(discount_total) if is_first else 0,
                    "payment":  payment,
                })
                order_subtotal += unit_price * output_qty
                is_first        = False
                success_count  += 1

            if has_error:
                continue

            is_cod = "貨到付款" in payment

            # 貨到付款補手續費
            if is_cod:
                cod_prod  = sku_map.get(COD_FEE_SKU)
                cod_name  = cod_prod.get("品名") if cod_prod else "黑貓貨到付款手續費"
                xlsx_rows.append({
                    "order_id": order_id,
                    "rcv_name": rcv_name,
                    "address":  address,
                    "rcv_phone": rcv_phone,
                    "sku":      COD_FEE_SKU,
                    "name":     cod_name,
                    "qty":      1,
                    "price":    COD_FEE_PRICE,
                    "discount": 0,
                    "payment":  payment,
                })

            cod_amount = int(order_subtotal - discount_total + COD_FEE_PRICE) if is_cod else ""
            csv_orders.append({
                "rcv_name":  rcv_name,
                "rcv_phone": rcv_phone,
                "address":   address,
                "order_id":  order_id,
                "cod":       cod_amount,
            })

        out_xlsx = _write_xlsx(self.output_dir.parent, self.source_type,
                               date_obj, f"{mmdd}.xlsx", xlsx_rows)
        out_csv  = _write_csv(self.output_dir.parent, self.source_type,
                              date_obj, f"{mmdd}黑貓.csv", csv_orders, date_obj)

        status = "completed" if fail_count == 0 else ("partial" if success_count > 0 else "failed")
        return ConversionResult(
            source_type   = self.source_type,
            success_count = success_count,
            fail_count    = fail_count,
            order_date    = order_date,
            output_files  = [out_xlsx, out_csv],
            errors        = errors,
            status        = status,
        )


# ── 輸出 ──────────────────────────────────────────────────────────────────

def _write_xlsx(base_dir: Path, source_type: str, date_obj: datetime,
                filename: str, rows: list[dict]) -> Path:
    p = output_path(base_dir, source_type, date_obj, filename)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(OUTPUT_HEADER)
    for r in rows:
        ws.append([
            r["order_id"], r["rcv_name"], r["address"], r["rcv_phone"],
            None,
            r["sku"], r["name"], r["qty"], r["price"],
            0, r["discount"], 0,
            None, None,
            r["payment"],
        ])
    wb.save(p)
    return p


def _write_csv(base_dir: Path, source_type: str, date_obj: datetime,
               filename: str, orders: list[dict], ref_date: datetime) -> Path:
    p = output_path(base_dir, source_type, date_obj, filename)
    ship_date = ref_date.strftime("%Y/%m/%d")
    next_date = (ref_date + timedelta(days=1)).strftime("%Y/%m/%d")
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(CSV_HEADER)
        for o in orders:
            w.writerow([
                o["rcv_name"], "", o["rcv_phone"], o["address"],
                o["cod"], 1, 1, "",
                o["order_id"],
                SENDER["deliver"],
                ship_date, next_date,
                SENDER["temp"], SENDER["size"],
                SENDER["name"], SENDER["tel"], SENDER["mobile"], SENDER["address"],
                "", "", "", "", "", "", "", "", "",
            ])
    return p


# ── 工具 ─────────────────────────────────────────────────────────────────

def _split(val) -> list[str]:
    if val is None:
        return []
    return [s.strip() for s in str(val).split("\n") if s.strip()]


def _num(val) -> float:
    try:
        return float(str(val).strip())
    except (ValueError, TypeError):
        return 0.0


def _parse_address(raw: str) -> str:
    """
    解析多行地址：郵遞區號\\n城市 行政區\\n街道\\n姓名
    僅保留「城市行政區 + 街道」。
    """
    if not raw:
        return ""
    lines = [l.strip() for l in raw.split("\n") if l.strip()]

    # 英文地址直接保留（排除純數字郵遞區號開頭的中文地址）
    if lines and re.match(r"^[A-Za-z]", lines[0]):
        return " ".join(lines)

    if len(lines) < 3:
        return "".join(lines)

    # line[0] = 郵遞區號（skip）；line[-1] = 姓名（skip）
    city_line = lines[1]
    street    = lines[2]

    # 判斷順序：若第一個詞不含「市/縣」但第二個詞含「市/縣」→ 行政區在前 → 對調
    parts = city_line.split()
    if len(parts) >= 2 and re.search(r"[市縣]$", parts[1]) and not re.search(r"[市縣]", parts[0]):
        city_dist = parts[1] + parts[0]
    else:
        city_dist = "".join(parts)

    # 街道已含完整城市行政區 → 直接使用
    if re.match(r"^.+[市縣].+[區鄉鎮市]", street):
        return street

    return city_dist + street
